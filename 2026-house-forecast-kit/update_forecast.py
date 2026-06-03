#!/usr/bin/env python3
"""
update_forecast.py  —  refresh the 2026 House forecast from house-polling.xlsx

Usage:
    python update_forecast.py            # reads ./house-polling.xlsx
    python update_forecast.py my.xlsx

What it does:
  * reads the Inputs sheet (generic ballot, redistricting shift, source, date)
  * reads the Districts sheet (optional per-district margin overrides)
  * rewrites data.js and rebuilds 2026-house-forecast.html (the shareable file)
  * stamps "Data updated" with the sheet's date, or today's date if blank
No internet or build tools required — just Python + openpyxl.
"""
import json, re, sys, math, datetime
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "house-polling.xlsx"

def load_data():
    raw = (HERE / "data.js").read_text()
    body = re.sub(r";\s*$", "", raw[len("window.HOUSE_DATA = "):])
    return json.loads(body)

def fmt_date(v):
    if v in (None, ""):
        v = datetime.date.today()
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%B %-d, %Y")
    return str(v).strip()

def read_inputs(ws):
    out = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        label = str(row[0]).lower()
        val = row[1] if len(row) > 1 else None
        if val in (None, ""):
            continue
        if "generic ballot" in label and "source" not in label:
            try: out["genericBallot"] = max(0.0, min(10.0, float(val)))
            except ValueError: pass
        elif "redistrict" in label or "shift" in label:
            try: out["redistrictShift"] = float(val)
            except ValueError: pass
        elif "source" in label:
            out["gbSource"] = str(val)
        elif "updated" in label or "date" in label:
            out["dataUpdated"] = fmt_date(val)
    return out

def read_districts(ws, code_to_id):
    overrides = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        code = str(row[0]).strip().upper()
        if code not in code_to_id:
            continue
        try:
            overrides[code_to_id[code]] = round(float(row[1]), 2)
        except (TypeError, ValueError):
            pass
    return overrides

def main():
    if not xlsx.exists():
        sys.exit(f"Spreadsheet not found: {xlsx}")
    data = load_data()
    fc, meta = data["forecast"], data["meta"]
    code_to_id = {f["code"].upper(): i for i, f in fc.items()}

    wb = load_workbook(xlsx, data_only=True)
    inputs = read_inputs(wb["Inputs"] if "Inputs" in wb.sheetnames else wb[wb.sheetnames[0]])
    meta.update(inputs)
    if "dataUpdated" not in inputs:
        meta["dataUpdated"] = fmt_date(None)  # today

    n_over = 0
    if "Districts" in wb.sheetnames:
        ov = read_districts(wb["Districts"], code_to_id)
        for did, m in ov.items():
            fc[did]["margin"] = m
        n_over = len(ov)

    # recompute baseline seat tally at the chosen generic ballot (for meta tidiness)
    anchor, gb = 6.5, meta.get("genericBallot", 6.5)
    swing = gb - anchor
    margins = sorted((f["margin"] + swing for f in fc.values()), reverse=True)
    dem = sum(1 for m in margins if m >= 0)
    meta["demSeats"], meta["gopSeats"] = dem, len(margins) - dem
    meta["leader"] = "D" if dem >= meta["gopSeats"] else "R"

    # write data.js
    (HERE / "data.js").write_text("window.HOUSE_DATA = " + json.dumps(data) + ";\n")

    # rebuild the standalone, shareable HTML
    html = (HERE / "index.html").read_text()
    css = (HERE / "styles.css").read_text()
    djs = (HERE / "data.js").read_text()
    ajs = (HERE / "app.js").read_text()
    html = html.replace('  <link rel="stylesheet" href="styles.css" />', "  <style>\n" + css + "\n  </style>")
    html = html.replace('  <script src="data.js"></script>', "  <script>\n" + djs + "\n  </script>")
    html = html.replace('  <script src="app.js"></script>', "  <script>\n" + ajs + "\n  </script>")
    (HERE / "2026-house-forecast.html").write_text(html)

    print("Updated 2026-house-forecast.html")
    print(f"  generic ballot   : D +{gb}")
    print(f"  redistrict shift : {meta.get('redistrictShift')}")
    print(f"  data updated     : {meta['dataUpdated']}")
    print(f"  district edits   : {n_over}")
    print(f"  baseline tally   : D {dem} / R {meta['gopSeats']}")

if __name__ == "__main__":
    main()
