import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

raw = open('data.js').read()
data = json.loads(re.sub(r';\s*$', '', raw[len('window.HOUSE_DATA = '):]))
fc, meta = data['forecast'], data['meta']

ARIAL = 'Arial'
BLUE = Font(name=ARIAL, color='0000FF')          # editable inputs
BOLD = Font(name=ARIAL, bold=True)
NORM = Font(name=ARIAL)
GREY = Font(name=ARIAL, color='808080', size=9)
HEADER = Font(name=ARIAL, bold=True, color='FFFFFF')
YELLOW = PatternFill('solid', fgColor='FFF6CC')
HEADFILL = PatternFill('solid', fgColor='1A1C20')
thin = Side(style='thin', color='DDDDDD')
BORDER = Border(bottom=thin)

wb = Workbook()

# ---------------- Inputs ----------------
s = wb.active; s.title = 'Inputs'
s.column_dimensions['A'].width = 42
s.column_dimensions['B'].width = 26
s.column_dimensions['C'].width = 60

s['A1'] = '2026 House Forecast — Inputs'; s['A1'].font = Font(name=ARIAL, bold=True, size=14)
s['A2'] = 'Edit the blue cells, save, then load this file on the page (or run update_forecast.py).'
s['A2'].font = GREY
r = 4
s.cell(r, 1, 'Setting').font = HEADER; s.cell(r, 2, 'Value').font = HEADER; s.cell(r, 3, 'Notes').font = HEADER
for c in range(1, 4):
    s.cell(r, c).fill = HEADFILL; s.cell(r, c).alignment = Alignment(horizontal='left')

rows = [
    ('Generic ballot (Dem margin, points)', meta.get('genericBallot', 6.5),
     'Democratic lead in the national House popular vote. 0–10. Sets the slider start.'),
    ('Redistricting shift — Plus (points)', meta.get('redistrictShift', -3.3),
     'Baseline shift for the Plus model. Negative = toward Republicans.'),
    ('Generic ballot source', meta.get('gbSource', 'Nate Silver · Silver Bulletin'),
     'Shown in the small-print note under the slider.'),
    ('Data updated (date)', meta.get('dataUpdated', 'June 3, 2026'),
     'Auto-stamps the “Model run · …” line. Leave blank to use the file’s save date.'),
]
for i, (label, val, note) in enumerate(rows):
    rr = 5 + i
    s.cell(rr, 1, label).font = NORM
    c = s.cell(rr, 2, val); c.font = BLUE; c.fill = YELLOW; c.alignment = Alignment(horizontal='left')
    s.cell(rr, 3, note).font = GREY
    for cc in range(1, 4):
        s.cell(rr, cc).border = BORDER

# ---------------- Districts ----------------
d = wb.create_sheet('Districts')
d.column_dimensions['A'].width = 12
d.column_dimensions['B'].width = 18
d.column_dimensions['C'].width = 8
d.column_dimensions['D'].width = 54

d['A1'] = 'District-level margins (optional fine-tuning)'; d['A1'].font = Font(name=ARIAL, bold=True, size=13)
d['A2'] = 'Margin = projected Democratic two-party margin at the D+6.5 baseline (negative = Republican). Edit any blue cell.'
d['A2'].font = GREY

hr = 4
for col, name in enumerate(['District', 'Margin (Dem +)', 'State', 'Note'], start=1):
    c = d.cell(hr, col, name); c.font = HEADER; c.fill = HEADFILL

rowsd = sorted(fc.values(), key=lambda x: x['code'])
for i, f in enumerate(rowsd):
    rr = hr + 1 + i
    d.cell(rr, 1, f['code']).font = NORM
    c = d.cell(rr, 2, round(f['margin'], 1)); c.font = BLUE; c.fill = YELLOW
    c.number_format = '0.0;-0.0'
    d.cell(rr, 3, f['state']).font = GREY
    for cc in range(1, 5):
        d.cell(rr, cc).border = BORDER

d.freeze_panes = 'A5'

wb.save('house-polling.xlsx')
print('wrote house-polling.xlsx with', len(rowsd), 'districts')
